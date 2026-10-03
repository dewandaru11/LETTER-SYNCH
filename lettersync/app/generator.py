from pathlib import Path
from docx import Document
from datetime import datetime
import subprocess
import sys


if getattr(sys, "frozen", False):
    BASE_DIR = Path(sys.executable).resolve().parent
    RESOURCE_DIR = Path(__file__).resolve().parent.parent
else:
    BASE_DIR = Path(__file__).resolve().parent.parent
    RESOURCE_DIR = BASE_DIR


TEMPLATE_DIR = RESOURCE_DIR / "templates"
OUTPUT_DIR = BASE_DIR / "output"

WORD_DIR = OUTPUT_DIR / "word"
PDF_DIR = OUTPUT_DIR / "pdf"



def _replace_in_paragraph(paragraph, replacements):

    for old, new in replacements.items():

        new = str(new or "")

        for run in paragraph.runs:

            if old in run.text:
                run.text = run.text.replace(
                    old,
                    new
                )

        full_text = paragraph.text

        if old in full_text:

            new_text = full_text.replace(
                old,
                new
            )

            if paragraph.runs:

                paragraph.runs[0].text = new_text

                for run in paragraph.runs[1:]:
                    run.text = ""


def _replace_everywhere(doc, replacements):

    for paragraph in doc.paragraphs:

        _replace_in_paragraph(
            paragraph,
            replacements
        )

    for table in doc.tables:

        for row in table.rows:

            for cell in row.cells:

                for paragraph in cell.paragraphs:

                    _replace_in_paragraph(
                        paragraph,
                        replacements
                    )

    for section in doc.sections:

        for paragraph in section.header.paragraphs:

            _replace_in_paragraph(
                paragraph,
                replacements
            )

        for table in section.header.tables:

            for row in table.rows:

                for cell in row.cells:

                    for paragraph in cell.paragraphs:

                        _replace_in_paragraph(
                            paragraph,
                            replacements
                        )

        for paragraph in section.footer.paragraphs:

            _replace_in_paragraph(
                paragraph,
                replacements
            )

        for table in section.footer.tables:

            for row in table.rows:

                for cell in row.cells:

                    for paragraph in cell.paragraphs:

                        _replace_in_paragraph(
                            paragraph,
                            replacements
                        )


def _replace_cell_text(cell, new_text):

    new_text = str(new_text or "")

    if not cell.paragraphs:

        cell.text = new_text
        return

    paragraph = cell.paragraphs[0]

    if paragraph.runs:

        paragraph.runs[0].text = new_text

        for run in paragraph.runs[1:]:

            run.text = ""

    else:

        paragraph.add_run(new_text)


def _cell_text(cell):

    return " ".join(
        line.strip()
        for line in cell.text.splitlines()
        if line.strip()
    )


def _find_table_with_header(doc):

    required = [
        "No",
        "Hari/Tanggal",
        "Jam Mulai Lembur",
        "Jam Selesai Lembur",
        "Jumlah Jam Lembur",
        "Laporan Lembur",
    ]

    for table in doc.tables:

        for row_index, row in enumerate(table.rows):

            texts = [
                _cell_text(cell).lower()
                for cell in row.cells
            ]

            combined = " | ".join(texts)

            score = 0

            for item in required:

                if item.lower() in combined:
                    score += 1

            if score >= 4:

                return table, row_index

    return None, None


def _remove_row(row):

    tbl = row._tr

    tbl.getparent().remove(tbl)



def _set_lembur_table_borders(table):
    """Memastikan semua garis tabel lembur tampil di Word."""

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl = table._tbl
    tblPr = tbl.tblPr

    # Hapus konfigurasi border lama jika ada.
    old_borders = tblPr.find(qn("w:tblBorders"))

    if old_borders is not None:
        tblPr.remove(old_borders)

    tbl_borders = OxmlElement("w:tblBorders")

    for edge in (
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV",
    ):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")
        tbl_borders.append(element)

    tblPr.append(tbl_borders)

    # Pastikan setiap cell juga mempunyai border.
    for row in table.rows:
        for cell in row.cells:

            tcPr = cell._tc.get_or_add_tcPr()

            old_cell_borders = tcPr.find(qn("w:tcBorders"))

            if old_cell_borders is not None:
                tcPr.remove(old_cell_borders)

            cell_borders = OxmlElement("w:tcBorders")

            for edge in (
                "top",
                "left",
                "bottom",
                "right",
            ):
                element = OxmlElement(f"w:{edge}")
                element.set(qn("w:val"), "single")
                element.set(qn("w:sz"), "6")
                element.set(qn("w:space"), "0")
                element.set(qn("w:color"), "000000")
                cell_borders.append(element)

            tcPr.append(cell_borders)


def _replace_lembur_template(doc, values):
    """
    Mengisi template Daftar Hadir Kerja Lembur.

    Mendukung:
    - format lama: satu tanggal
    - format baru: banyak tanggal melalui values["LEMBUR_ITEMS"]
    """

    import re
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    nama = str(values.get("NAMA", "") or "")
    nip = str(values.get("NIP", "") or "")
    unit_kerja = str(values.get("UNIT_KERJA", "") or "")
    tanggal_surat = str(
        values.get("TANGGAL_SURAT")
        or values.get("TANGGAL_PEMBUATAN_SURAT")
        or ""
    )

    # ---------------------------------------------------------
    # DATA MULTI-TANGGAL
    # ---------------------------------------------------------
    items = values.get("LEMBUR_ITEMS") or []

    # Kompatibilitas dengan format lama
    if not items:
        items = [{
            "tanggal": str(values.get("TANGGAL_LEMBUR") or values.get("TANGGAL") or ""),
            "tanggal_iso": str(values.get("TANGGAL_ISO") or ""),
            "jam_mulai": str(values.get("JAM_MULAI") or ""),
            "jam_selesai": str(values.get("JAM_SELESAI") or ""),
            "total": str(values.get("TOTAL_JAM") or ""),
            "laporan": str(values.get("LAPORAN_LEMBUR") or ""),
        }]

    def replace_runs(paragraph, replacements):
        for run in paragraph.runs:
            text = run.text
            for old, new in replacements:
                if old in text:
                    text = text.replace(old, str(new))
            run.text = text

    def set_cell_border(cell):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()

        tcBorders = tcPr.first_child_found_in("w:tcBorders")
        if tcBorders is None:
            tcBorders = OxmlElement("w:tcBorders")
            tcPr.append(tcBorders)

        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            tag = "w:" + edge
            element = tcBorders.find(qn(tag))

            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)

            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), "6")
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), "000000")

    def set_cell_width(cell, width):
        tcPr = cell._tc.get_or_add_tcPr()
        tcW = tcPr.first_child_found_in("w:tcW")

        if tcW is None:
            tcW = OxmlElement("w:tcW")
            tcPr.append(tcW)

        tcW.set(qn("w:w"), str(width))
        tcW.set(qn("w:type"), "dxa")

    def set_cell_text(cell, text):
        cell.text = str(text or "")

        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.underline = False

    # ---------------------------------------------------------
    # GANTI INFORMASI HEADER
    # ---------------------------------------------------------
    for paragraph in doc.paragraphs:

        if "REKAM MEDIS" in paragraph.text and unit_kerja:
            replace_runs(
                paragraph,
                [("REKAM MEDIS", unit_kerja)]
            )

        if "Hariyadi Adha" in paragraph.text and nama:
            replace_runs(
                paragraph,
                [
                    ("Hariyadi Adha, S.Pd", nama),
                    ("Hariyadi Adha", nama),
                ]
            )

        if "07012102016113030" in paragraph.text and nip:
            replace_runs(
                paragraph,
                [("07012102016113030", nip)]
            )

        # TANGGAL PEMBUATAN SURAT
        # Pada template, "Barabai" dan tanggal berada pada beberapa RUN.
        # Ubah hanya RUN yang berisi "Barabai" agar format/alignment tetap terjaga.
        if tanggal_surat and "Barabai" in paragraph.text:
            if "Mengetahui" in paragraph.text:
                for i, run in enumerate(paragraph.runs):
                    if run.text.strip() == "Barabai":
                        run.text = f"Barabai, {tanggal_surat}"

                        # Hapus RUN tanggal lama setelah "Barabai".
                        for extra_run in paragraph.runs[i + 1:]:
                            extra_run.text = ""

                        break


        if (
            "NIP. 197307171998032008" in paragraph.text
            and "NIPK. 07012102016113030" in paragraph.text
            and nip
        ):
            replace_runs(
                paragraph,
                [("07012102016113030", nip)]
            )

    # ---------------------------------------------------------
    # CARI TABEL UTAMA
    # ---------------------------------------------------------
    target_table = None

    for table in doc.tables:

        if not table.rows:
            continue

        header = " ".join(
            cell.text.strip()
            for cell in table.rows[0].cells
        )

        if (
            "Hari/Tanggal" in header
            and "Jam Mulai Lembur" in header
            and "Jam Selesai Lembur" in header
            and "Jumlah Jam Lembur" in header
            and "Laporan Lembur" in header
        ):
            target_table = table
            break

    if target_table is None:
        raise ValueError(
            "Tabel lembur utama tidak ditemukan."
        )

    # ---------------------------------------------------------
    # HAPUS TABEL LAIN
    # ---------------------------------------------------------
    for table in list(doc.tables):

        if table._tbl is target_table._tbl:
            continue

        tbl = table._tbl
        parent = tbl.getparent()

        if parent is not None:
            parent.remove(tbl)

    # ---------------------------------------------------------
    # HAPUS SEMUA ROW DATA, SISAKAN HEADER
    # ---------------------------------------------------------
    tbl = target_table._tbl

    while len(target_table.rows) > 1:
        tbl.remove(target_table.rows[-1]._tr)

    # ---------------------------------------------------------
    # HEADER
    # ---------------------------------------------------------
    headers = [
        "No",
        "Hari/Tanggal",
        "Jam Mulai Lembur",
        "Jam Selesai Lembur",
        "Jumlah Jam Lembur",
        "Laporan Lembur",
    ]

    header_row = target_table.rows[0]

    if len(header_row.cells) != 6:
        raise ValueError(
            f"Tabel lembur membutuhkan 6 kolom, "
            f"tetapi ditemukan {len(header_row.cells)}."
        )

    widths = [
        650,
        1450,
        1250,
        1250,
        1150,
        3250,
    ]

    for index, text in enumerate(headers):
        set_cell_text(header_row.cells[index], text)
        set_cell_width(header_row.cells[index], widths[index])
        set_cell_border(header_row.cells[index])

    # ---------------------------------------------------------
    # BUAT SEMUA BARIS LEMBUR
    # ---------------------------------------------------------
    for index, item in enumerate(items, start=1):

        row = target_table.add_row()

        if len(row.cells) != 6:
            raise ValueError(
                "Baris tabel lembur harus memiliki 6 kolom."
            )

        data = [
            str(index),
            str(item.get("tanggal", "") or ""),
            str(item.get("jam_mulai", "") or ""),
            str(item.get("jam_selesai", "") or ""),
            str(item.get("total", "") or ""),
            str(item.get("laporan", "") or ""),
        ]

        for col, text in enumerate(data):
            set_cell_text(row.cells[col], text)
            set_cell_width(row.cells[col], widths[col])
            set_cell_border(row.cells[col])

    # ---------------------------------------------------------
    # HAPUS UNDERLINE DARI SELURUH DOKUMEN
    # ---------------------------------------------------------
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.underline = False

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.underline = False

    return doc

def render_template(
    template_name,
    output_name,
    values
):

    template = TEMPLATE_DIR / template_name

    if not template.exists():

        raise FileNotFoundError(
            f"Template tidak ditemukan:\n{template}"
        )

    WORD_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    output = WORD_DIR / output_name

    doc = Document(template)

    # ========================================================
    # TEMPLATE LEMBUR
    # ========================================================

    if template_name.lower() == (
        "daftar hadir kerja lembur.docx"
    ):

        _replace_lembur_template(
            doc,
            values
        )

    else:

        _replace_everywhere(
            doc,
            values
        )

    doc.save(output)

    return output


def find_soffice():

    possible_paths = [

        Path(
            r"C:\Program Files\LibreOffice\program\soffice.exe"
        ),

        Path(
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"
        ),

    ]

    for path in possible_paths:

        if path.exists():

            return path

    return None


def create_pdf_from_docx(docx_path):

    docx_path = Path(docx_path)

    if not docx_path.exists():

        return None

    soffice = find_soffice()

    if soffice is None:

        return None

    PDF_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    pdf_path = (
        PDF_DIR /
        f"{docx_path.stem}.pdf"
    )

    try:

        subprocess.run(
            [
                str(soffice),
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(PDF_DIR),
                str(docx_path)
            ],
            capture_output=True,
            text=True,
            timeout=120
        )

        if pdf_path.exists():

            return pdf_path

        return None

    except Exception:

        return None


def today_id():

    return datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )
